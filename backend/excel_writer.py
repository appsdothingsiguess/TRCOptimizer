import datetime

import openpyxl


def write_intake_row(
    iiq_ticket: str,
    asset_tag: str,
    serial: str,
    school_name: str,
    tech_initials: str,
    break_count: int,
    excel_path: str,
) -> None:
    wb = openpyxl.load_workbook(excel_path)
    ws = wb["Sheet1"]

    target_row = None
    for row_num in range(4, ws.max_row + 2):
        if ws.cell(row=row_num, column=2).value is None:
            target_row = row_num
            break

    if target_row is None:
        wb.close()
        raise ValueError("No empty intake row found in Excel sheet")

    ws.cell(row=target_row, column=2).value = datetime.date.today()
    ws.cell(row=target_row, column=3).value = iiq_ticket
    ws.cell(row=target_row, column=5).value = asset_tag
    ws.cell(row=target_row, column=6).value = serial
    ws.cell(row=target_row, column=7).value = school_name
    ws.cell(row=target_row, column=8).value = tech_initials
    ws.cell(row=target_row, column=9).value = int(break_count)

    wb.save(excel_path)
    wb.close()
